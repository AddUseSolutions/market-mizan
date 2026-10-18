#!/usr/bin/env python3
"""Apply reviewer Excel translations (EN/AM/OM) into frontend i18n files."""
from __future__ import annotations

import json
import re
import sys
from pathlib import Path

try:
    import openpyxl
except ImportError:
    sys.exit("openpyxl required")

ROOT = Path(__file__).resolve().parents[1]
XLSX = Path.home() / "Downloads" / "MMizan_Website_Translations_EN_AM_OM.xlsx"
LC = ROOT / "frontend" / "src" / "context" / "LanguageContext.jsx"
EXTRA = ROOT / "frontend" / "src" / "i18n" / "extraLanguages.js"

# Manual AM strings as unicode escapes (ASCII-safe source).
AM = {
    "area": "\u12a0\u12ab\u1263\u1262",
    "showing": "\u1260\u12a0\u1295\u12f5 \u1308\u133d {n}",
    "page": "\u1308\u133d {page} \u12a8{total}",
    "mm": "\u121b\u122d\u12ac\u1275 \u121a\u12db\u1295",
    "contact": "\u12a0\u12f5\u122b\u123b\u1366",
    "available": "\u12ed\u1308\u129b\u120d",
    "active": "\u1295\u1241",
    "rented": "\u1270\u12a8\u122b\u12ed\u1277\u120d",
    "vacant": "\u1263\u12f6",
    "sold": "\u1270\u123d\u1327\u120d",
    "listing_mode": "\u12e8\u121b\u1235\u1273\u12c8\u1242\u12eb \u12d3\u12ed\u1290\u1275",
    "availability": "\u1270\u1308\u129d\u1290\u1275",
    "updated": "\u1208\u1218\u1328\u1228\u123b \u130a\u12dc \u12e8\u1270\u123b\u123b\u1208\u12cd\u1366 \u121a\u12eb\u12dd\u12eb 2026",
    "eth_suffix": "\u1363 \u12a2\u1275\u12ee\u1335\u12eb",
    "for_rent": "\u1208\u12aa\u122b\u12ed",
    "for_sale": "\u1208\u123d\u12eb\u132d",
}

EN_TO_KEYS: dict[str, list[str] | None] = {
    "Analyze with MIZAN and find the right property": ["heroTitle"],
    "The trusted portal to compare and navigate the real estate market in Ethiopia": ["heroSub"],
    "Upload your listing": ["heroUploadCta"],
    "Buy": ["buy", "searchBuy"],
    "Rent": ["rent", "searchRent"],
    "Type": ["searchType"],
    "Area": None,
    "Bedrooms": ["searchBedrooms", "bedrooms", "beds"],
    "Search": ["searchSubmit"],
    "Apartment / Condo": ["typeApartmentCondo"],
    "Villa / House": ["typeVillaHouse"],
    "Studio": ["typeStudio"],
    "Room": ["typeRoom"],
    "Office": ["typeOffice"],
    "Shop / Retail": ["typeShopRetail"],
    "Commercial Building": ["typeCommercialBuilding"],
    "Hotel": ["typeHotel"],
    "Warehouse": ["typeWarehouse"],
    "Residential Land": ["typeResidentialLand"],
    "Commercial Land": ["typeCommercialLand"],
    "Agricultural Land": ["typeAgriculturalLand"],
    "listings": ["listingsCount"],
    "Sort": ["sort"],
    "Recommended": ["sortRecommended"],
    "Newest": ["sortLatest"],
    "Price: low to high": ["sortPriceAsc"],
    "Price: high to low": ["sortPriceDesc"],
    "Size": ["sortSize"],
    "Living area": ["livingArea"],
    "Furnished": ["furnishedLabel"],
    "View details": ["viewDetails"],
    "Compare": ["compareMode", "compareShort"],
    "For rent": ["searchForRent"],
    "For Sale": ["forSale"],
    "Verified": ["verified"],
    "Your property aggregator for Addis Ababa: clear listings, one calm place to browse.": ["footerLead"],
    "Addis Ababa, Ethiopia": ["footerLocation"],
    "Explore": ["footerExplore"],
    "Listings": ["navListings"],
    "Map": ["navMap"],
    "List your property": ["footerListYourProperty"],
    "Contact": ["footerContact"],
    "Login": ["navLogin"],
    "Sign in": ["signIn"],
    "Menu": ["menu"],
    "Properties": ["properties"],
    "Legal": ["footerLegal", "legalEyebrow"],
    "Privacy policy": ["footerPrivacyPolicy", "privacyTitle"],
    "Terms of use": ["footerTermsOfUse", "termsTitle"],
    "Legal notice (Imprint)": ["footerLegalNotice", "legalNoticeTitle"],
    "Privacy": ["footerPrivacyShort"],
    "Terms": ["footerTermsShort"],
    "Imprint": ["footerImprint"],
    "Back to listings": ["backToListings"],
    "monthly rent": ["monthlyRentSuffix"],
    "Features": ["detailFeatures"],
    "No features listed.": ["detailNoFeatures"],
    "Location": ["detailLocation"],
    "Mentioned in description:": ["detailLocationMentioned"],
    "Request removal of this listing": ["detailRequestRemoval", "removalTitle"],
    "Your bridge to trusted property experts": ["holisticServicesTitle"],
    "Choose what you need help with and we'll connect you with the right specialist.": [
        "holisticServicesLead"
    ],
    "Property management": ["supplierPropertyManagement"],
    "Tenant support, maintenance coordination, rent handling": ["supplierPropertyManagementDesc"],
    "Interior design": ["supplierInteriorDesign"],
    "Staging, furnishing, renovation guidance": ["supplierInteriorDesignDesc"],
    "Craftsmanship": ["supplierCraftsmanship"],
    "Repairs, finishing, and build-out support": ["supplierCraftsmanshipDesc"],
    "Mortgage / bank": ["supplierMortgage"],
    "Financing support and bank introductions": ["supplierMortgageDesc"],
    "Insurance": ["supplierInsurance"],
    "Property coverage and risk protection": ["supplierInsuranceDesc"],
    "Agreement": ["termsAgreementTitle"],
    "By accessing or using Market Mizan (mmizan.com), you agree to these terms. If you do not agree, please do not use the service.": [
        "termsAgreementBody"
    ],
    "Service description": ["termsServiceTitle"],
    "Market Mizan aggregates publicly available property listings from third-party sources for information purposes. We do not guarantee completeness, accuracy, or availability of listings. Always verify details with the listing provider before making decisions.": [
        "termsServiceBody"
    ],
    "Accounts": ["termsAccountsTitle"],
    "Where we offer accounts, you are responsible for keeping your credentials confidential and for all activity under your account. Notify us promptly of any unauthorised use.": [
        "termsAccountsBody"
    ],
    "Acceptable use": ["termsAcceptableTitle"],
    "You must not misuse the site (e.g. scraping in violation of our rules, introducing malware, overloading systems, or using the service for unlawful purposes). We may suspend access if necessary.": [
        "termsAcceptableBody"
    ],
    "Disclaimer": ["termsDisclaimerTitle"],
    'The service is provided "as is". To the extent permitted by law, we disclaim warranties and limit liability arising from your use of aggregated information or third-party content.': [
        "termsDisclaimerBody"
    ],
    "Changes": ["termsChangesTitle"],
    "We may update these terms from time to time. Continued use after changes constitutes acceptance of the revised terms.": [
        "termsChangesBody"
    ],
    "Who we are": ["privacyWhoTitle"],
    'Market Mizan ("we", "us") operates the website mmizan.com and related services. This page explains how we handle personal data when you use our site.': [
        "privacyWhoBody"
    ],
    "Data we collect": ["privacyDataTitle"],
    "Depending on how you use Market Mizan, we may process: account details (e.g. name, email) if you register; messages you send via contact forms; and technical data such as IP address, browser type, and usage logs required to run and secure the service.": [
        "privacyDataBody"
    ],
    "How we use data": ["privacyUseTitle"],
    "We use data to provide and improve the platform, respond to enquiries, authenticate users where applicable, and meet legal obligations. We do not sell your personal data.": [
        "privacyUseBody"
    ],
    "Third parties": ["privacyThirdTitle"],
    "We may use hosting, analytics, or authentication providers (for example Google Analytics 4 and sign-in with Google). Those providers process data under their own terms and only as needed to deliver and improve the service. Analytics helps us understand which pages are used; it is not used to sell your data.": [
        "privacyThirdBody"
    ],
    "Your rights": ["privacyRightsTitle"],
    "Information pursuant to applicable transparency rules.": ["legalNoticeLead"],
    "Operator": ["legalOperatorTitle"],
    "Responsible for content": ["legalResponsibleTitle"],
    "The operator named above is responsible for own content on these pages under applicable law. Liability for links to third-party websites is limited as provided by law; we do not control external sites.": [
        "legalResponsibleBody"
    ],
    "Intellectual property": ["legalIpTitle"],
    "Texts, graphics, and branding on Market Mizan are protected where applicable. Listing data may originate from third parties; respective rights remain with the original sources.": [
        "legalIpBody"
    ],
}


def esc(s: str) -> str:
    return '"' + s.replace("\\", "\\\\").replace('"', '\\"').replace("\n", "\\n") + '"'


def replace_key(block: str, key: str, value: str) -> tuple[str, int]:
    pat = re.compile(rf"(\n\s*{re.escape(key)}:\s*)(?:\"(?:\\.|[^\"\\])*\"|'(?:\\.|[^'\\])*')")
    return pat.subn(lambda m: m.group(1) + esc(value), block, count=1)


def ensure_keys(block: str, keys: dict[str, str], indent: str = "    ") -> str:
    for k, v in keys.items():
        if re.search(rf"\n\s*{re.escape(k)}:", block):
            block, _ = replace_key(block, k, v)
        else:
            trimmed = block.rstrip()
            if not trimmed.endswith(","):
                trimmed += ","
            block = trimmed + f"\n{indent}{k}: {esc(v)}\n"
    return block


def fix_email(s: str | None) -> str | None:
    if not s:
        return s
    return s.replace("hello@add-use.ch", "hello@mmizan.com")


def load_from_xlsx() -> dict[str, dict[str, str]]:
    if not XLSX.exists():
        sys.exit(f"Missing Excel: {XLSX}")
    wb = openpyxl.load_workbook(XLSX, data_only=True)
    ws = wb["Translations"]
    by_key: dict[str, dict[str, str]] = {}

    for row in ws.iter_rows(min_row=5, values_only=True):
        cells = list(row) + [None] * 6
        _num, _section, en, am, om, _notes = cells[:6]
        if not isinstance(en, str) or not en.strip():
            continue
        en = en.strip()
        if "Please replace the above" in en:
            continue
        if en.startswith("Last updated: April"):
            by_key["privacyUpdated"] = {
                "en": "Last updated: April 2026",
                "am": AM["updated"],
                "om": "Yeroo dhumaaf kan haaromfame: Ebla 2026",
            }
            by_key["termsUpdated"] = dict(by_key["privacyUpdated"])
            continue
        if en == "Last updated:":
            continue
        if en.startswith("Where applicable law grants you rights"):
            by_key["privacyRightsBody"] = {
                "en": fix_email(en) or en,
                "am": fix_email(am) or "",
                "om": fix_email(om) or "",
            }
            continue
        if en not in EN_TO_KEYS:
            continue
        keys = EN_TO_KEYS[en]
        if not keys:
            continue
        for k in keys:
            by_key[k] = {"en": en, "am": am or "", "om": om or ""}

    kt = {
        r[0]: (r[1], r[2])
        for r in wb["Key terms"].iter_rows(min_row=4, values_only=True)
        if r[0]
    }
    addis = kt.get("Addis Ababa", (AM["area"], "Finfinnee"))
    # Prefer key-term Addis; fall back
    addis_am = addis[0] if addis[0] else "\u12a0\u12f2\u1235 \u12a0\u1260\u1263"
    addis_om = addis[1] if addis[1] else "Finfinnee"

    for_rent = by_key.get("searchForRent", {})
    by_key["forRent"] = {
        "en": "For Rent",
        "am": for_rent.get("am") or AM["for_rent"],
        "om": for_rent.get("om") or "Kiraaf",
    }
    for_sale = by_key.get("forSale", {})
    by_key["searchForSale"] = {
        "en": "For sale",
        "am": for_sale.get("am") or AM["for_sale"],
        "om": for_sale.get("om") or "Gurgurtaaf",
    }
    by_key["searchArea"] = {"en": "Area", "am": AM["area"], "om": "Naannoo"}
    by_key["showingPerPage"] = {
        "en": "Showing {n} per page",
        "am": AM["showing"],
        "om": "Fuula tokkotti {n}",
    }
    by_key["pageOf"] = {
        "en": "Page {page} of {total}",
        "am": AM["page"],
        "om": "Fuula {page} kan {total} keessaa",
    }
    by_key["legalOperatorName"] = {
        "en": "Market Mizan",
        "am": AM["mm"],
        "om": "Maarkeet Miizaan",
    }
    by_key["legalOperatorAddress"] = {
        "en": "Addis Ababa, Ethiopia",
        "am": f"{addis_am}{AM['eth_suffix']}",
        "om": f"{addis_om}, Itoophiyaa",
    }
    by_key["legalContactLabel"] = {
        "en": "Contact:",
        "am": AM["contact"],
        "om": "Quunnamtii:",
    }
    by_key["statusAvailable"] = {
        "en": "Available",
        "am": AM["available"],
        "om": "Ni argama",
    }
    by_key["statusActive"] = {
        "en": "Active",
        "am": AM["active"],
        "om": "Kan hojjetaa jiru",
    }
    by_key["statusRented"] = {
        "en": "Rented",
        "am": AM["rented"],
        "om": "Kan kireeffame",
    }
    by_key["statusVacant"] = {
        "en": "Vacant",
        "am": AM["vacant"],
        "om": "Kan duwwaa ta'e",
    }
    by_key["statusSold"] = {"en": "Sold", "am": AM["sold"], "om": "Kan gurgurame"}
    by_key["detailListingMode"] = {
        "en": "Listing type",
        "am": AM["listing_mode"],
        "om": "Gosa beeksisaa",
    }
    by_key["detailAvailability"] = {
        "en": "Availability",
        "am": AM["availability"],
        "om": "Argamuu",
    }
    return by_key


def patch_language_context(by_key: dict[str, dict[str, str]]) -> None:
    text = LC.read_text(encoding="utf-8")
    m_en = re.search(r"(\n  en: \{)([\s\S]*?)(\n  \},\n  am: \{)", text)
    m_am = re.search(r"(\n  am: \{)([\s\S]*?)(\n  \}\n\};)", text)
    if not m_en or not m_am:
        sys.exit("Could not find en/am blocks in LanguageContext.jsx")

    en_block, am_block = m_en.group(2), m_am.group(2)
    en_extra: dict[str, str] = {}
    am_extra: dict[str, str] = {}
    for key, vals in by_key.items():
        en_block2, n1 = replace_key(en_block, key, vals["en"])
        am_block2, n2 = replace_key(am_block, key, vals["am"])
        if n1:
            en_block = en_block2
        else:
            en_extra[key] = vals["en"]
        if n2:
            am_block = am_block2
        else:
            am_extra[key] = vals["am"]

    en_block = ensure_keys(en_block, en_extra)
    am_block = ensure_keys(am_block, am_extra)
    text = (
        text[: m_en.start(2)]
        + en_block
        + text[m_en.end(2) : m_am.start(2)]
        + am_block
        + text[m_am.end(2) :]
    )
    LC.write_text(text, encoding="utf-8")


def patch_om(by_key: dict[str, dict[str, str]]) -> None:
    text = EXTRA.read_text(encoding="utf-8")
    m = re.search(r"(export const OM_STRINGS = \{)([\s\S]*?)(\n\};)", text)
    if not m:
        sys.exit("OM_STRINGS not found")
    block = m.group(2)
    extra: dict[str, str] = {}
    for key, vals in by_key.items():
        block2, n = replace_key(block, key, vals["om"])
        if n:
            block = block2
        else:
            extra[key] = vals["om"]
    block = ensure_keys(block, extra, indent="  ")
    text = text[: m.start(2)] + block + text[m.end(2) :]
    EXTRA.write_text(text, encoding="utf-8")


def main() -> None:
    by_key = load_from_xlsx()
    (ROOT / "scripts" / "_i18n_patch.json").write_text(
        json.dumps(by_key, ensure_ascii=False, indent=2), encoding="utf-8"
    )
    patch_language_context(by_key)
    patch_om(by_key)
    print(f"Applied {len(by_key)} keys from {XLSX.name}")


if __name__ == "__main__":
    main()
